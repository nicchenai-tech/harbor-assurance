"""Read formats into locatable source lines without executing document content."""
from pathlib import Path
import hashlib
import re
import json
import subprocess
import sys
import os

MAX_BYTES = 20 * 1024 * 1024


def safe_path(root, name):
    root = Path(root).resolve()
    path = (root / name).resolve()
    if not path.is_relative_to(root):
        raise ValueError('Attachment path leaves the input directory.')
    return path


def read_document(root, name, timeout=60):
    """Kill a stalled parser process; a broken source cannot stall the inbox."""
    safe_path(root, name)
    env=os.environ.copy()
    package_root=str(Path(__file__).resolve().parent.parent)
    env['PYTHONPATH']=package_root+os.pathsep+env.get('PYTHONPATH','')
    try:
        completed=subprocess.run([sys.executable,'-m','harbor.parsers',str(root),name],
                                 capture_output=True,text=True,timeout=timeout,env=env)
        if completed.returncode:
            raise ValueError('Document worker failed: '+completed.stderr[-300:])
        return json.loads(completed.stdout)
    except subprocess.TimeoutExpired:
        return {'file':name,'format':Path(name).suffix.lstrip('.'),'sha256':None,
                'lines':[],'errors':[f'Parser timed out after {timeout:g} seconds. Replace or retry this source.'],
                'scan':False,'method':'native_text'}


def _read_document(root, name):
    p = safe_path(root, name)
    result = {'file':name,'format':p.suffix.lower().lstrip('.'),'sha256':None,
              'lines':[],'errors':[],'scan':False,'method':'native_text'}
    if not p.exists():
        result['errors'] = ['Attachment file is missing.']; return result
    if p.stat().st_size > MAX_BYTES:
        result['errors'] = ['Attachment exceeds the 20 MB limit.']; return result
    data = p.read_bytes()
    result['sha256'] = hashlib.sha256(data).hexdigest()
    if not data:
        result['errors'] = ['Attachment is empty.']; return result
    def add(text, locator):
        if str(text).strip():
            result['lines'].append({'text':str(text).strip(), 'locator':locator})
    try:
        if p.suffix.lower() == '.txt':
            for i, line in enumerate(data.decode('utf-8-sig').splitlines(),1):
                add(line, {'line':i})
        elif p.suffix.lower() == '.pdf':
            from pypdf import PdfReader
            reader = PdfReader(p)
            if len(reader.pages)>50: raise ValueError('PDF exceeds 50-page limit.')
            result['pages'] = len(reader.pages)
            for n,page in enumerate(reader.pages,1):
                text = page.extract_text() or ''
                if not text.strip():
                    result['scan'] = True
                    try:
                        from .ocr import recognize_pdf_page
                        ocr_lines=recognize_pdf_page(p,n-1)
                        for i,item in enumerate(ocr_lines,1):
                            add(item['text'],{'page':n,'line':i,'bbox':item['bbox']})
                        backend=ocr_lines[0].get('backend','ocr') if ocr_lines else 'ocr'
                        result['method']=backend+'_ocr' if not backend.endswith('_ocr') else backend
                    except Exception as exc:
                        result['errors'].append(f'OCR unavailable or failed: {str(exc)[:200]}')
                for i,line in enumerate(text.splitlines(),1):
                    add(line, {'page':n,'line':i})
            if result['scan'] and not result['lines']:
                result['errors'].append('No readable text. Replace the file or enter values with source evidence.')
        elif p.suffix.lower() == '.docx':
            from docx import Document
            doc = Document(p)
            for i,para in enumerate(doc.paragraphs,1): add(para.text, {'paragraph':i})
            for t,table in enumerate(doc.tables,1):
                for r,row in enumerate(table.rows,1):
                    cells=[cell.text for cell in row.cells]
                    if cells: add(cells[0]+': '+'\n'.join(cells[1:]), {'table':t,'row':r})
        elif p.suffix.lower() == '.xlsx':
            from openpyxl import load_workbook
            book = load_workbook(p, data_only=False, read_only=True)
            for sheet in book:
                if sheet.max_row>10000 or sheet.max_column>100: raise ValueError('Spreadsheet exceeds supported dimensions.')
                for row in sheet:
                    cells=[c for c in row if c.value is not None]
                    if not cells: continue
                    label=str(cells[0].value)
                    values=[str(c.value) for c in cells[1:]]
                    if any(c.data_type == 'f' for c in cells):
                        result['errors'].append('Formula value requires review; formulas are not executed.')
                    add(label+(': '+'\n'.join(values) if values else ''),
                        {'sheet':sheet.title,'cell':cells[-1].coordinate,'label_cell':cells[0].coordinate})
            book.close()
        else:
            result['errors'].append('Unsupported attachment format.')
    except Exception as exc:
        result['errors'].append(f'{type(exc).__name__}: {str(exc)[:220]}')
    return result


def document_role(document):
    header = '\n'.join(l['text'] for l in document['lines'][:7]).upper()
    # Negative types first. Filename never participates.
    if re.search(r'COMMERCIAL INVOICE|PACKING LIST|CERTIFICATE OF ORIGIN',header): return 'OTHER'
    if re.search(r'SHIPPING INSTRUCTION|BILL OF LADING INSTRUCTION|B/?L\.? INSTRUCTION',header): return 'SI'
    if re.search(r'BILL OF LADING|DRAFT B/?L',header): return 'BL'
    return 'UNKNOWN'


if __name__=='__main__':
    print(json.dumps(_read_document(sys.argv[1],sys.argv[2])))
