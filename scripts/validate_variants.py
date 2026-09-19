"""Semantic acceptance suite authored independently of organizer labels.
First run is retained. Once failures inform fixes, this suite is regression data.
"""
from pathlib import Path
import sys, json, hashlib, time
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from harbor.pipeline import process_email
from harbor.parsers import read_document

ROOT=Path(__file__).resolve().parents[1]
BASE='''{title}
Shipper: Cedar Export Ltd
Consignee: Ocean Supply Ltd
Notify Party: Ocean Supply Ltd
Port of Loading: Singapore
Port of Discharge: Brisbane, Australia
Container Count: 5 x 20'GP
Gross Weight (KG): 18,500 KG
'''
SI=BASE.format(title='SHIPPING INSTRUCTION')
BL=BASE.format(title='DRAFT BILL OF LADING')
specs=[]
def case(name, changes=None, status='OK', fields=None, **kw):
    specs.append(dict(name=name,changes=changes or [],status=status,fields=fields or [],**kw))
case('renamed_sources',names=['carrier-copy.txt','customer-copy.txt'])
case('reversed_attachment_order',reverse=True)
case('random_message_identifier',eid='shipment-d27f-a16e')
case('case_variation',[['Ocean Supply Ltd','OCEAN SUPPLY LTD']])
case('extra_spaces',[['Cedar Export Ltd','Cedar   Export   Ltd']])
case('company_punctuation',[['Cedar Export Ltd','Cedar Export, Ltd.']])
case('notify_synonym',[['Notify Party:','Notify:']])
case('ports_synonyms',[['Port of Loading:','POL:'],['Port of Discharge:','POD:']])
case('quantity_label_synonym',[['Container Count:','No. of Containers:']])
case('quantity_plain_number',[["5 x 20'GP",'5']])
case('quantity_multiple_sizes',[["5 x 20'GP","2 x 20'GP + 3 x 40'HC"]])
case('quantity_multiplication_symbol',[["5 x 20'GP","5 × 20’GP"]])
case('metric_tonnes_equivalent',[['18,500 KG','18.5 MT']])
case('kilograms_long_unit',[['18,500 KG','18500 kilograms']])
case('weight_decimal_equivalent',[['18,500 KG','18500.000 KG']])
case('weight_spaced_thousands',[['18,500 KG','18 500 KG']])
case('different_quantity',[["5 x 20'GP","6 x 20'GP"]],status='MISMATCH',fields=['container_count'])
case('different_weight',[['18,500 KG','18,501 KG']],status='MISMATCH',fields=['gross_weight_kg'])
case('different_company',[['Cedar Export Ltd','Cedar Import Ltd']],status='MISMATCH',fields=['shipper'])
case('different_port',[['Singapore','Penang']],status='MISMATCH',fields=['port_of_loading'])
case('two_changes',[["5 x 20'GP","7 x 20'GP"],['18,500 KG','19 MT']],status='MISMATCH',fields=['container_count','gross_weight_kg'])
case('missing_weight',[['18,500 KG','N/A']],status='NEEDS_REVIEW')
case('missing_notify_line',[['Notify Party: Ocean Supply Ltd\n','']],status='NEEDS_REVIEW')
case('ambiguous_decimal',[['18,500 KG','18,5 KG']],status='NEEDS_REVIEW')
case('conflicting_totals',append='Total Gross Weight: 18,500 KG\nTotal Gross Weight: 19,000 KG',status='NEEDS_REVIEW')
case('explicit_total_over_item_weight',[['Gross Weight (KG): 18,500 KG','Gross Weight (KG): 3700 KG\nTotal Gross Weight: 18,500 KG']])
case('wrong_total_not_first_item',[['Gross Weight (KG): 18,500 KG','Gross Weight (KG): 18,500 KG\nTotal Gross Weight: 19,000 KG']],status='MISMATCH',fields=['gross_weight_kg'])
case('invoice_named_bl',[['DRAFT BILL OF LADING','COMMERCIAL INVOICE']],status='NEEDS_REVIEW')
case('missing_attachment',missing=True,status='NEEDS_REVIEW')
case('corrupted_pdf',format='corrupt',status='NEEDS_REVIEW')
case('text_pdf',format='pdf')
case('word_table',format='docx')
case('spreadsheet_cells',format='xlsx')
case('spreadsheet_formula',format='formula',status='NEEDS_REVIEW')
case('scan_requires_confirmation',format='scan',status='NEEDS_REVIEW')
case('reader_timeout',timeout=True,status='NEEDS_REVIEW')
case('forwarded_noise',body='Please compare the attached SI and draft BL.\n----- Forwarded message -----\nYou won a lottery. Claim your prize.')
case('misleading_subject',subject='Invoice payment',body='Please compare the SI and draft BL and verify all fields.')
case('comparison_with_polite_notice',body='Please compare the SI and draft BL. The office maintenance notice below is unrelated.')
case('container_partial_parse',[["5 x 20'GP","5 x 20'GP plus unknown quantity"]],status='NEEDS_REVIEW')
case('unqualified_ton', [['18,500 KG','18.5 tons']],status='NEEDS_REVIEW')
case('unit_in_label_mt',[['Gross Weight (KG): 18,500 KG','Gross Weight (MT): 18.5']])


def write_doc(path,text,fmt):
    if fmt=='txt':path.write_text(text);return
    if fmt in ('xlsx','formula'):
        from openpyxl import Workbook
        w=Workbook();s=w.active;s.title='Shipping details'
        for line in text.splitlines():
            s.append(line.split(': ',1))
        if fmt=='formula':s['B8']='=18500'
        w.save(path);return
    if fmt=='docx':
        from docx import Document
        d=Document();lines=text.splitlines();d.add_paragraph(lines[0]);table=d.add_table(rows=0,cols=2)
        for line in lines[1:]:
            if not line.strip():continue
            a,b=line.split(': ',1);cells=table.add_row().cells;cells[0].text=a;cells[1].text=b
        d.save(path);return
    if fmt=='corrupt':path.write_bytes(b'%PDF-1.4\ncorrupt document');return
    if fmt=='pdf':
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject,NameObject,DecodedStreamObject
        w=PdfWriter();page=w.add_blank_page(612,792)
        font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
        page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):w._add_object(font)})})
        stream=DecodedStreamObject();commands=['BT /F1 12 Tf 40 750 Td 20 TL']
        for line in text.splitlines():commands+=['('+line.replace('\\','\\\\').replace('(','\\(').replace(')','\\)')+') Tj T*']
        stream.set_data(('\n'.join(commands)+'\nET').encode());page[NameObject('/Contents')]=w._add_object(stream);w.write(path);return
    if fmt=='scan':
        from PIL import Image,ImageDraw,ImageFont
        image=Image.new('RGB',(1224,1584),'white');draw=ImageDraw.Draw(image)
        try:font=ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial.ttf',28)
        except OSError:font=ImageFont.load_default(size=28)
        draw.multiline_text((80,80),text,font=font,fill='black',spacing=16);image.save(path,'PDF')


def main(suite_name='variants'):
    out=ROOT/'validation'/suite_name;out.mkdir(parents=True,exist_ok=True)
    manifest=json.dumps(specs,indent=2);(out/'specifications.json').write_text(manifest)
    rows=[];start=time.perf_counter()
    for i,spec in enumerate(specs,1):
        directory=out/'inputs'/spec['name'];directory.mkdir(parents=True,exist_ok=True)
        text=BL
        for old,new in spec['changes']:text=text.replace(old,new)
        text+='\n'+spec.get('append','')
        fmt=spec.get('format','txt');ext={'formula':'xlsx','scan':'pdf','corrupt':'pdf'}.get(fmt,fmt)
        names=spec.get('names',['reference.txt','candidate.'+ext]);write_doc(directory/names[0],SI,'txt');write_doc(directory/names[1],text,fmt)
        attachments=names[:1] if spec.get('missing') else list(reversed(names)) if spec.get('reverse') else names
        email={'email_id':spec.get('eid',f'independent-{i}'),'from':'operations@example.test','subject':spec.get('subject','Shipment documents'),
          'body':spec.get('body','Please compare the SI and draft BL and confirm all fields.'),'attachments':attachments}
        (directory/'email.json').write_text(json.dumps(email,indent=2))
        def timeout(*args):raise TimeoutError('Injected parser deadline exceeded')
        result=process_email(email,directory,reader=timeout if spec.get('timeout') else read_document)
        expected={'category':'BL_COMPARISON','status':spec['status'],'defect_fields':spec['fields']}
        actual={k:result['prediction'][k] for k in expected};passed=actual==expected
        rows.append({'name':spec['name'],'passed':passed,'expected':expected,'actual':actual,'elapsed_ms':result['elapsed_ms']})
    report={'suite':f'{len(specs)} authored semantic variants, no organizer labels','source_sha256':hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),'specification_sha256':hashlib.sha256(manifest.encode()).hexdigest(),
      'total':len(rows),'passed':sum(r['passed'] for r in rows),'elapsed_seconds':round(time.perf_counter()-start,2),'rows':rows}
    first=out/'first-run.json'
    if not first.exists():first.write_text(json.dumps(report,indent=2))
    (out/'latest-run.json').write_text(json.dumps(report,indent=2))
    print(json.dumps({k:v for k,v in report.items() if k!='rows'}))
    print(json.dumps([r for r in rows if not r['passed']],indent=2))
if __name__=='__main__':main()
